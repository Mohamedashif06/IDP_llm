"""Direct file reader for Word (.docx) and Excel (.xlsx) files.

Extracts text, key-value pairs, and tables deterministically without VLM.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Tuple


def classify_text_by_keywords(text: str) -> Tuple[str, float]:
    """Assign document type using keyword rules over extracted digital text."""
    lower = text.lower()
    if any(k in lower for k in ["invoice", "tax id", "bill no", "subtotal", "gstin", "vendor"]):
        return "invoice", 0.95
    if any(k in lower for k in ["certificate", "degree", "awarded to", "grade", "conferred"]):
        return "certificate", 0.95
    if any(k in lower for k in ["roll no", "register no", "mark sheet", "end semester", "internal assessment"]):
        return "mark_sheet", 0.95
    if any(k in lower for k in ["identity card", "id no", "date of birth", "valid until", "employee id"]):
        return "id_card", 0.95
    if any(k in lower for k in ["attendance", "register", "present", "absent"]):
        return "attendance_register", 0.95
    return "other", 0.60


def read_docx(source: str | bytes) -> Dict[str, Any]:
    """Read digital Word document directly with python-docx."""
    import docx

    if isinstance(source, bytes):
        doc = docx.Document(io.BytesIO(source))
    else:
        doc = docx.Document(source)

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)

    fields: Dict[str, Any] = {}
    key_value_pairs: List[Dict[str, Any]] = []

    for p in paragraphs:
        match = re.match(r"^\s*([^:=]{2,50})\s*[:=]\s*(.+)\s*$", p)
        if match:
            k = re.sub(r"\s+", "_", match.group(1).strip().lower())
            v = match.group(2).strip()
            fields[k] = v
            key_value_pairs.append({"key": match.group(1).strip(), "value": v, "confidence": 1.0})

    tables_data: List[Dict[str, Any]] = []
    for idx, table in enumerate(doc.tables, start=1):
        table_rows = []
        headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
        for row in table.rows[1:]:
            row_vals = [cell.text.strip() for cell in row.cells]
            table_rows.append(row_vals)
        tables_data.append({
            "table_name": f"Table_{idx}",
            "headers": headers,
            "rows": table_rows
        })

    doc_type, type_conf = classify_text_by_keywords(full_text)

    return {
        "source": getattr(source, "name", "digital_word.docx"),
        "page_count": 1,
        "engine": "python-docx (Direct Read)",
        "document_type": doc_type,
        "classification_confidence": type_conf,
        "reliability": "high (direct read)",
        "confidence": 1.0,
        "fields": fields,
        "key_value_pairs": key_value_pairs,
        "tables": tables_data,
        "text": full_text,
        "needs_review": [],
        "pages": []
    }


def read_xlsx(source: str | bytes) -> Dict[str, Any]:
    """Read digital Excel workbook directly with openpyxl."""
    import openpyxl

    if isinstance(source, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(source), data_only=True)
    else:
        wb = openpyxl.load_workbook(source, data_only=True)

    fields: Dict[str, Any] = {}
    tables_data: List[Dict[str, Any]] = []
    text_lines: List[str] = []

    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        headers = []
        sheet_rows = []
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            row_cells = [str(c).strip() if c is not None else "" for c in row]
            if not any(row_cells):
                continue
            text_lines.append(" | ".join(row_cells))
            if not headers:
                headers = row_cells
            else:
                sheet_rows.append(row_cells)
                if len(row_cells) >= 2 and row_cells[0] and row_cells[1]:
                    k = re.sub(r"\s+", "_", row_cells[0].lower())
                    if len(k) < 40:
                        fields[k] = row_cells[1]

        tables_data.append({
            "table_name": sheetname,
            "headers": headers,
            "rows": sheet_rows
        })

    full_text = "\n".join(text_lines)
    doc_type, type_conf = classify_text_by_keywords(full_text)

    return {
        "source": getattr(source, "name", "digital_excel.xlsx"),
        "page_count": len(wb.sheetnames),
        "engine": "openpyxl (Direct Read)",
        "document_type": doc_type,
        "classification_confidence": type_conf,
        "reliability": "high (direct read)",
        "confidence": 1.0,
        "fields": fields,
        "tables": tables_data,
        "text": full_text,
        "needs_review": [],
        "pages": []
    }
